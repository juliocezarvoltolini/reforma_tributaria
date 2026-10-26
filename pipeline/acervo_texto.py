"""Camada de texto do acervo: um Markdown por ato, derivado do original arquivado.

Para um LLM (ou uma pessoa) ler o acervo sem extrair o original na hora:
`indices/acervo_texto/<id>.md`, com front-matter de procedência (URL, arquivo,
sha256 do original) e o texto:

- PDF: por página (`## Página N`), para a citação levar a página. Tabela de PDF
  sai com colunas possivelmente embaralhadas — o aviso vai no arquivo;
- pacote de schemas (ZIP de XSD): a TABELA DOS CAMPOS de IBS, CBS e IS lida dos
  próprios XSD (caminho, tipo, ocorrência, tamanho, enumeração, padrão,
  descrição, arquivo) — a fonte sem ambiguidade do leiaute;
- planilha: uma tabela por aba; HTML e JSON da RFB: o texto.

Derivado e regenerável (`python pipeline/acervo_texto.py`); nunca editar à mão.
A validação confere, por hash, que cada arquivo veio do original atual e que o
corpo não foi mexido (`sha256_corpo`) — sem reextrair tudo a cada commit.

Campos do leiaute: entra toda subárvore cujo elemento ou tipo tenha "IBS" ou
"CBS" no nome, ou que se chame IS/gIS/ISTot (Imposto Seletivo). ISSQN e ISUF
não são Imposto Seletivo e ficam fora.
"""
from __future__ import annotations

import io
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import yaml

import acervo
from comum import DIR_INDICES, MANIFEST, caminho_fonte, gravar_texto, sha256_bytes

DIR_TEXTO = DIR_INDICES / "acervo_texto"
# Sobe a cada mudança no que o gerador produz: a validação acusa o texto antigo.
VERSAO_GERADOR = 3
XS = "{http://www.w3.org/2001/XMLSchema}"
AVISO_PDF = ("Texto extraído do PDF (pypdf), página a página. Tabelas podem sair com as colunas "
             "embaralhadas: para campos do leiaute, use o texto do pacote de schemas da mesma versão; "
             "para o resto, confira no original.")


# --------------------------------------------------------------------- XSD

def _local(nome: str | None) -> str:
    return (nome or "").split(":")[-1]


def _do_imposto(nome: str) -> bool:
    return bool(re.search(r"IBS|CBS", nome) or re.fullmatch(r"g?IS(Tot)?", nome))


def _restricoes(simples: ET.Element | None) -> dict:
    """Tamanho, enumeração e padrão de um xs:simpleType (com a base)."""
    if simples is None:
        return {}
    r = simples.find(f"{XS}restriction")
    if r is None:
        return {}
    val = {f.tag[len(XS):]: f.get("value") for f in r}
    tamanho = val.get("length") or (
        f"{val.get('minLength', '')}-{val.get('maxLength', '')}" if ("minLength" in val or "maxLength" in val)
        else "")
    if not tamanho and "totalDigits" in val:
        tamanho = f"{val['totalDigits']},{val.get('fractionDigits', '0')}"
    return {"base": r.get("base", ""), "tamanho": tamanho,
            "enumeracao": ", ".join(e.get("value") for e in r.findall(f"{XS}enumeration")),
            "padrao": " | ".join(p.get("value") for p in r.findall(f"{XS}pattern"))}


def campos_do_leiaute(pacote: bytes) -> list[dict]:
    """Grupos e campos de IBS/CBS/IS de um pacote de schemas, com o caminho a partir
    da raiz do documento, na ordem do leiaute."""
    complexos, simples, elementos = {}, {}, {}
    raizes: list = []  # todo elemento de topo de todo XSD (eventos repetem o nome)

    def ler(conteudo: bytes, prefixo: str) -> None:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            for nome in z.namelist():
                if nome.lower().endswith(".zip"):  # pacote dentro do pacote
                    ler(z.read(nome), f"{prefixo}{nome}/")
                    continue
                if not nome.lower().endswith(".xsd") or "xmldsig" in nome.lower():
                    continue
                raiz = ET.fromstring(z.read(nome))  # respeita o encoding declarado
                for filho in raiz:
                    n = filho.get("name")
                    if filho.tag == f"{XS}complexType":
                        complexos[n] = (filho, prefixo + nome)
                    elif filho.tag == f"{XS}simpleType":
                        simples[n] = (filho, prefixo + nome)
                    elif filho.tag == f"{XS}element":
                        elementos.setdefault(n, (filho, prefixo + nome))
                        raizes.append((filho, prefixo + nome))

    ler(pacote, "")
    campos: list[dict] = []
    vistos: set[tuple[str, str]] = set()  # (XSD da raiz, caminho)
    alcancados: set[str] = set()  # tipos complexos que algum elemento usa

    def _particulas(no: ET.Element, arquivo: str, pilha: tuple, escolha: bool = False):
        """(elemento, XSD, dentro de xs:choice?) dos filhos de um tipo complexo."""
        for filho in no:
            if filho.tag == f"{XS}element":
                yield filho, arquivo, escolha
            elif filho.tag in (f"{XS}sequence", f"{XS}choice", f"{XS}all"):
                yield from _particulas(filho, arquivo, pilha, escolha or filho.tag == f"{XS}choice")
            elif filho.tag in (f"{XS}complexContent", f"{XS}simpleContent"):
                ext = filho.find(f"{XS}extension")
                if ext is not None:
                    base = _local(ext.get("base"))
                    if base in complexos and base not in pilha:
                        yield from _particulas(complexos[base][0], complexos[base][1], pilha + (base,), escolha)
                    yield from _particulas(ext, arquivo, pilha, escolha)

    def visitar(el: ET.Element, arquivo: str, caminho: str, incluir: bool, pilha: tuple, raiz: str,
                escolha: bool = False) -> None:
        if el.get("ref"):
            ref = _local(el.get("ref"))
            if ref not in elementos:
                return
            alvo, arq_alvo = elementos[ref]
            el = ET.Element(alvo.tag, {**alvo.attrib, **{k: v for k, v in el.attrib.items() if k != "ref"}})
            el.extend(list(alvo))
            arquivo = arq_alvo
        nome = el.get("name", "")
        tipo = _local(el.get("type"))
        if caminho and tipo in tipos_de_raiz:  # envelope: o documento é raiz própria
            return
        caminho = f"{caminho}/{nome}" if caminho else nome
        incluir = incluir or _do_imposto(nome) or bool(tipo and _do_imposto(tipo)) or tipo == "TIS"
        complexo = el.find(f"{XS}complexType")
        if tipo:
            alcancados.add(tipo)
        if complexo is None and tipo in complexos and tipo not in pilha:
            complexo, arq_tipo = complexos[tipo]
        else:
            arq_tipo = arquivo
        if incluir and (raiz, caminho) not in vistos:
            vistos.add((raiz, caminho))
            rest = _restricoes(el.find(f"{XS}simpleType")) or (
                _restricoes(simples[tipo][0]) if tipo in simples else {})
            doc = el.find(f"{XS}annotation/{XS}documentation")
            maximo = el.get("maxOccurs", "1")
            campos.append({
                "caminho": caminho,
                "tipo": el.get("type") and tipo or rest.get("base") or ("(grupo)" if complexo is not None else ""),
                "ocorrencia": f"{el.get('minOccurs', '1')}-{'N' if maximo == 'unbounded' else maximo}",
                "tamanho": rest.get("tamanho", ""), "enumeracao": rest.get("enumeracao", ""),
                "padrao": rest.get("padrao", ""),
                "descricao": " ".join((doc.text or "").split()) if doc is not None else "",
                "arquivo": arquivo, "escolha": "sim" if escolha else ""})
        if complexo is not None and len(pilha) < 40:
            for filho, arq, em_escolha in _particulas(complexo, arq_tipo, pilha + ((tipo,) if tipo else ())):
                visitar(filho, arq, caminho, incluir, pilha + ((tipo,) if tipo else ()), raiz, em_escolha)

    tipos_de_raiz = {_local(el.get("type")) for el, _ in raizes if el.get("type")}
    for el, arquivo in raizes:
        visitar(el, arquivo, "", False, (), arquivo)
    # Pacote sem nenhum elemento de topo (schema de evento: só tipos, a raiz está em
    # outro XSD): cada tipo complexo que nenhum elemento alcança vira raiz, e o
    # caminho começa pelo nome dele. Com elementos de topo, não — tipos de envelope
    # (retorno, processamento) repetiriam todos os campos com outro prefixo.
    for nome, (complexo, arquivo) in list(complexos.items()):
        if elementos or nome in alcancados:
            continue
        for filho, arq, em_escolha in _particulas(complexo, arquivo, (nome,)):
            visitar(filho, arq, nome, _do_imposto(nome), (nome,), arquivo, em_escolha)
    return campos


def _celula(valor) -> str:
    return " ".join(str(valor if valor is not None else "").split()).replace("|", "\\|")


def tabela_markdown(cabecalho: list[str], linhas: list[list]) -> str:
    saida = ["| " + " | ".join(cabecalho) + " |", "|" + "---|" * len(cabecalho)]
    saida += ["| " + " | ".join(_celula(c) for c in linha) + " |" for linha in linhas]
    return "\n".join(saida)


# --------------------------------------------------------------- Markdown

def _corpo(dados: bytes, ext: str) -> tuple[str, str | None]:
    """(corpo, aviso) do texto derivado de um original."""
    if ext == "pdf":
        paginas = acervo.paginas_do_pdf(dados)
        return ("\n\n".join(f"## Página {i}\n\n{t.strip()}" for i, t in enumerate(paginas, 1)), AVISO_PDF)
    if ext == "zip":
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            nomes = z.namelist()
        campos = campos_do_leiaute(dados)
        partes = ["## Arquivos do pacote\n\n" + "\n".join(f"- {n}" for n in nomes)]
        if campos:
            # `escolha`: o elemento está num xs:choice — a ocorrência só vale se a
            # alternativa for a escolhida. Última coluna, para não mudar as outras.
            colunas = ["caminho", "tipo", "ocorrencia", "tamanho", "enumeracao", "padrao", "descricao", "arquivo",
                       "escolha"]
            partes.append("## Campos de IBS, CBS e IS (lidos dos XSD)\n\n"
                          + tabela_markdown(colunas, [[c[k] for k in colunas] for c in campos]))
        else:
            partes.append("## Campos de IBS, CBS e IS\n\nNenhum grupo ou campo de IBS, CBS ou IS nos XSD "
                          "do pacote (só nomes e conteúdo acima).")
        return "\n\n".join(partes), "Tabela lida mecanicamente dos XSD do pacote (pipeline/acervo_texto.py)."
    if ext == "xlsx":
        from openpyxl import load_workbook

        livro = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        partes = []
        for aba in livro.worksheets:
            linhas = [list(r) for r in aba.iter_rows(values_only=True) if any(c is not None for c in r)]
            if linhas:
                largura = max(len(r) for r in linhas)
                linhas = [r + [None] * (largura - len(r)) for r in linhas]
                partes.append(f"## Planilha: {aba.title}\n\n"
                              + tabela_markdown([_celula(c) or "-" for c in linhas[0]], linhas[1:]))
        return "\n\n".join(partes), None
    return acervo.texto_bruto(dados, ext).strip(), None


def markdown_do_ato(reg: dict, registro: dict, dados: bytes) -> str:
    corpo, aviso = _corpo(dados, Path(registro["arquivo"]).suffix[1:])
    # \r solto (JSON da RFB) vira \n na releitura do .md: normaliza antes do hash.
    corpo = corpo.replace("\r\n", "\n").replace("\r", "\n").strip() + "\n"
    meta = {"id": reg["id"], "titulo": reg.get("titulo"), "tipo": reg.get("tipo"),
            "versao": reg.get("versao"), "data": reg.get("data"),
            "documento_fiscal": reg.get("documento_fiscal"), "url": reg.get("url"),
            "arquivo_original": registro["arquivo"], "sha256_original": registro["sha256"],
            "gerado_por": "pipeline/acervo_texto.py", "versao_gerador": VERSAO_GERADOR, "aviso": aviso,
            "sha256_corpo": sha256_bytes(corpo.encode("utf-8"))}
    meta = {k: v for k, v in meta.items() if v is not None}
    return "---\n" + yaml.safe_dump(meta, allow_unicode=True, sort_keys=False, width=1000) + "---\n\n" + corpo


def _separar(md: str) -> tuple[dict, str]:
    _, cabeca, corpo = md.split("---\n", 2)
    return yaml.safe_load(cabeca) or {}, corpo[1:] if corpo.startswith("\n") else corpo


def problemas_do_texto(md: str, docs: dict) -> list[str]:
    meta, corpo = _separar(md)
    registro = docs.get(f"acervo/{meta.get('id')}")
    if registro is None:
        return [f"{meta.get('id')}: sem original no manifest"]
    probs = []
    if meta.get("sha256_original") != registro["sha256"]:
        probs.append(f"{meta['id']}: o original mudou desde a extração — rode pipeline/acervo_texto.py")
    if meta.get("versao_gerador") != VERSAO_GERADOR:
        probs.append(f"{meta['id']}: gerado por versão anterior do gerador — rode pipeline/acervo_texto.py")
    if meta.get("sha256_corpo") != sha256_bytes(corpo.encode("utf-8")):
        probs.append(f"{meta['id']}: corpo editado à mão (sha256_corpo não confere) — regenere")
    return probs


def problemas_da_camada(regs: list[dict], docs: dict, textos: dict[str, str]) -> list[str]:
    """Todo ato arquivado tem o seu texto, íntegro; nenhum texto sobra."""
    esperados = {r["id"] for r in regs if "bloqueio" not in r and f"acervo/{r['id']}" in docs}
    probs = [f"acervo_texto {rid}: sem texto — rode pipeline/acervo_texto.py"
             for rid in sorted(esperados - set(textos))]
    probs += [f"acervo_texto {rid}: texto de ato fora do acervo — rode pipeline/acervo_texto.py"
              for rid in sorted(set(textos) - esperados)]
    for rid in sorted(esperados & set(textos)):
        probs += [f"acervo_texto {p}" for p in problemas_do_texto(textos[rid], docs)]
    return probs


# ------------------------------------------------------------------ busca

def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto.lower()) if not unicodedata.combining(c))


def buscar_no_texto(textos: dict[str, str], termo: str) -> list[dict]:
    """Linhas que contêm o termo (sem acento nem caixa), com o ato, a versão e a página."""
    alvo = _sem_acento(termo)
    achados = []
    for rid, md in textos.items():
        meta, corpo = _separar(md)
        pagina = None
        for linha in corpo.split("\n"):
            m = re.match(r"## Página (\d+)$", linha)
            if m:
                pagina = int(m.group(1))
            elif alvo in _sem_acento(linha):
                achados.append({"id": rid, "versao": meta.get("versao"), "pagina": pagina,
                                "trecho": linha.strip()})
    return achados


def carregar_textos() -> dict[str, str]:
    return {p.stem: p.read_text(encoding="utf-8") for p in sorted(DIR_TEXTO.glob("*.md"))}


# ------------------------------------------------------------------ gerar

def main() -> int:
    import json

    docs = json.loads(MANIFEST.read_text(encoding="utf-8"))["documentos"]
    DIR_TEXTO.mkdir(parents=True, exist_ok=True)
    ids = set()
    for reg in acervo.carregar():
        registro = docs.get(f"acervo/{reg['id']}")
        if registro is None or "bloqueio" in reg:
            continue
        ids.add(reg["id"])
        destino = DIR_TEXTO / f"{reg['id']}.md"
        if destino.exists() and not problemas_do_texto(destino.read_text(encoding="utf-8"), docs):
            continue
        gravar_texto(destino, markdown_do_ato(reg, registro, caminho_fonte(registro).read_bytes()))
        print(f"  {destino.relative_to(DIR_INDICES.parent).as_posix()}")
    for velho in DIR_TEXTO.glob("*.md"):
        if velho.stem not in ids:
            velho.unlink()
            print(f"  removido {velho.name} (fora do acervo)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
